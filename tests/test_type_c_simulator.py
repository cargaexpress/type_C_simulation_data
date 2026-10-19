"""Unit tests for Type C File Simulator with Defect Injection (F5, Sec 5.2 and 5.3)."""

import csv
import re
from pathlib import Path

import openpyxl
from simulators.type_c.catalogs import get_carrier_config
from simulators.type_c.cli import main as cli_main
from simulators.type_c.generator import TypeCSimulator, compute_file_sha256


class TestTypeCSimulator:
    """Integration and unit test suite for Type C simulator."""

    def test_reproducibility_byte_to_byte_csv(self, tmp_path: Path):
        """Validate that two runs with the same seed produce exactly the same file (identical hash)."""
        sim1 = TypeCSimulator(carrier=1, seed=42, error_rate=0.1)
        sim2 = TypeCSimulator(carrier=1, seed=42, error_rate=0.1)

        file1 = sim1.generate_file(output_dir=tmp_path / "run1", num_rows=80, file_format="csv")
        file2 = sim2.generate_file(output_dir=tmp_path / "run2", num_rows=80, file_format="csv")

        hash1 = compute_file_sha256(file1)
        hash2 = compute_file_sha256(file2)

        assert hash1 == hash2, (
            "Executions with the same seed must produce the exact byte-to-byte SHA-256 hash"
        )
        assert file1.read_text(encoding="utf-8") == file2.read_text(encoding="utf-8")

    def test_reproducibility_different_seeds(self, tmp_path: Path):
        """Different seeds must produce different files."""
        sim_a = TypeCSimulator(carrier=1, seed=42, error_rate=0.1)
        sim_b = TypeCSimulator(carrier=1, seed=99, error_rate=0.1)

        file_a = sim_a.generate_file(output_dir=tmp_path / "run_a", num_rows=50, file_format="csv")
        file_b = sim_b.generate_file(output_dir=tmp_path / "run_b", num_rows=50, file_format="csv")

        assert compute_file_sha256(file_a) != compute_file_sha256(file_b)

    def test_clean_generation_without_errors(self, tmp_path: Path):
        """With error_rate=0, file is generated clean without repeated headers or summary rows."""
        sim = TypeCSimulator(carrier=1, seed=42, error_rate=0.0)
        csv_file = sim.generate_file(output_dir=tmp_path, num_rows=50, file_format="csv")

        with open(csv_file, encoding="utf-8") as f:
            reader = list(csv.reader(f))

        header = reader[0]
        data_rows = reader[1:]

        assert header == sim.config.columns
        assert len(data_rows) == 50

        # Must not contain repeated header rows
        for row in data_rows:
            assert row != header, "Must not have repeated headers in clean generation"
            assert not any("TOTAL REGISTROS" in str(col) for col in row), "Must not have total row"
            assert not any(str(col).startswith("FAKE_WAYBILL") for col in row), (
                "Must not have fake waybills"
            )

            # Validate standard ISO datetime format: YYYY-MM-DD HH:MM:SS
            ts_str = row[sim.config.columns.index(sim.config.timestamp_col)]
            assert re.match(r"^\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2}$", ts_str), (
                f"Date must match standard ISO: {ts_str}"
            )

    def test_corrupted_structure_injection(self, tmp_path: Path):
        """With error_rate > 0, resulting file contains repeated headers and summary footer."""
        sim = TypeCSimulator(carrier=1, seed=42, error_rate=0.2)
        csv_file = sim.generate_file(output_dir=tmp_path, num_rows=60, file_format="csv")

        with open(csv_file, encoding="utf-8") as f:
            content = f.read()

        # Verify summary footer row presence
        assert "TOTAL REGISTROS:" in content
        assert "--- FIN DE LOTE ---" in content

        # Verify presence of repeated headers in body
        lines = content.strip().split("\n")
        header_line = ",".join(sim.config.columns)
        assert lines[0] == header_line
        repeated_headers = [line for line in lines[1:] if line == header_line]
        assert len(repeated_headers) >= 1, "Must contain at least one repeated header in body"

    def test_mixed_date_formats_presence(self, tmp_path: Path):
        """With error_rate > 0, presence of mixed DD/MM/YYYY and MM/DD/YYYY formats must be detected."""
        sim = TypeCSimulator(carrier=1, seed=42, error_rate=0.3)
        events = sim.generate_events(num_rows=100)

        date_col = sim.config.timestamp_col
        timestamps = [row[date_col] for row in events if isinstance(row.get(date_col), str)]

        # Mixed pattern: DD/MM/YYYY HH:MM or MM/DD/YYYY HH:MM
        slash_dates = [ts for ts in timestamps if "/" in ts]
        assert len(slash_dates) > 0, "Must contain dates with interleaved slash format"

        # Validate pattern matches day/month/year hour:minute
        for dt_str in slash_dates:
            assert re.match(r"^\d{2}/\d{2}/\d{4} \d{2}:\d{2}$", dt_str), (
                f"Corrupted date must match format: {dt_str}"
            )

    def test_fake_guides_and_duplicates_injected(self):
        """Validate presence of fake waybills and duplicate records with error_rate > 0."""
        sim = TypeCSimulator(carrier=1, seed=42, error_rate=0.2)
        events = sim.generate_events(num_rows=80)

        # 1. Fake waybills
        fake_guides = [
            row for row in events if str(row.get("waybill_no", "")).startswith("FAKE_WAYBILL_")
        ]
        assert len(fake_guides) > 0, "Must have injected nonexistent waybills"

        # 2. Duplicate rows (exclude repeated headers or footers)
        tracking_values = [
            row.get("waybill_no") for row in events if "TOTAL" not in str(row.get("waybill_no"))
        ]
        assert len(tracking_values) > len(set(tracking_values)), (
            "Must have repeated/duplicate waybills in events"
        )

    def test_xlsx_format_generation(self, tmp_path: Path):
        """Validate Excel (.xlsx) file format generation."""
        sim = TypeCSimulator(carrier=2, seed=42, error_rate=0.1)
        xlsx_file = sim.generate_file(output_dir=tmp_path, num_rows=40, file_format="xlsx")

        assert xlsx_file.exists()
        assert xlsx_file.suffix == ".xlsx"

        wb = openpyxl.load_workbook(xlsx_file)
        assert sim.config.code in wb.sheetnames
        ws = wb[sim.config.code]

        rows = list(ws.iter_rows(values_only=True))
        assert len(rows) > 40
        assert list(rows[0]) == sim.config.columns

    def test_support_for_all_three_carriers(self, tmp_path: Path):
        """Verify that all 3 carriers can be simulated correctly with their catalogs."""
        for c_id in [1, 2, 3]:
            config = get_carrier_config(c_id)
            assert len(config.status_catalog) >= 6, (
                f"Carrier {c_id} must have at least 6 statuses in catalog"
            )

            sim = TypeCSimulator(carrier=c_id, seed=100 + c_id, error_rate=0.0)
            target_path = sim.generate_file(output_dir=tmp_path / f"c_{c_id}", num_rows=30)
            assert target_path.exists()

    def test_cli_execution(self, tmp_path: Path, capsys):
        """Verify CLI command execution with arguments."""
        out_dir = tmp_path / "cli_output"
        exit_code = cli_main(
            [
                "--seed",
                "42",
                "--carrier",
                "1",
                "--output-dir",
                str(out_dir),
                "--rows",
                "25",
                "--error-rate",
                "0.1",
                "--format",
                "csv",
            ]
        )

        assert exit_code == 0
        captured = capsys.readouterr()
        assert "TYPE C SIMULATOR - GENERATION COMPLETED SUCCESSFULLY" in captured.out
        assert "SHA-256 Hash" in captured.out
        assert any(out_dir.glob("*.csv"))
